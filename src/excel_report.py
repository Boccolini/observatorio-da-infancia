"""Generate the consolidated Excel deliverable."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .geo import REGIOES


HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _write_sheet(writer: pd.ExcelWriter, name: str, df: pd.DataFrame, freeze: str = "B2") -> None:
    df.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]
    ws.freeze_panes = freeze
    for col_idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        max_len = max(
            (len(str(v)) for v in df[col].astype(str).head(200)),
            default=10,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 32)


def write_excel(
    df_uf: pd.DataFrame,
    df_regiao: pd.DataFrame,
    sih_long: pd.DataFrame,
    sim_long: pd.DataFrame,
    sinasc: pd.DataFrame,
    pop: pd.DataFrame,
    metadata: dict[str, str],
    out_path: Path,
) -> Path:
    """Write the multi-sheet workbook used as the project deliverable."""
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        meta = pd.DataFrame(
            list(metadata.items()), columns=["Campo", "Valor"]
        )
        _write_sheet(writer, "00_Metadados", meta, freeze="A2")

        # 01 - UF level (full indicators)
        _write_sheet(writer, "01_UF_Indicadores", df_uf)

        # 02 - Regiao + Brasil
        _write_sheet(writer, "02_Regiao_Brasil", df_regiao)

        # 03 - SIH counts (ano x UF x faixa)
        _write_sheet(writer, "03_SIH_Internacoes", sih_long)

        # 04 - SIM counts
        _write_sheet(writer, "04_SIM_Obitos", sim_long)

        # 05 - SINASC NV
        _write_sheet(writer, "05_SINASC_NV", sinasc)

        # 06 - Population
        _write_sheet(writer, "06_Populacao_IBGE", pop)

        # 07 - Series temporais por regiao
        rates = ["taxa_int_menor_1_por_1000NV",
                 "taxa_obt_menor_1_por_1000NV",
                 "taxa_int_1a4_por_100k",
                 "taxa_obt_1a4_por_100k",
                 "taxa_int_menor_5_por_100k",
                 "taxa_obt_menor_5_por_100k",
                 "taxa_obt_menor_5_por_1000NV"]
        for rate in rates:
            wide = df_regiao.pivot_table(
                index="ano", columns="regiao", values=rate
            ).reindex(columns=REGIOES + ["Brasil"]).round(2)
            _write_sheet(writer, f"07_{rate[:25]}", wide.reset_index())
    return out_path
